"""合同审批审查系统 · 一键自测脚本（无需启动服务，直接验证代码可运行）。

运行：
  cd contract-review-system
  .venv/Scripts/python tests/test_run.py

脚本会：
  1) 自动生成一份「混合风险」测试合同 测试合同_混合样例.docx
  2) 走通 解析 -> 规则审查 -> 风险汇总 -> 回写文本 -> 报告 全链路
  3) 断言关键结果并打印 PASS/FAIL
该 .docx 也可在系统「审查大盘」用「上传并审查」手动上传验证。
"""
import os
import sys
import traceback

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document  # 生成测试合同用（已在 venv 安装）

from backend import store
from backend.models import BizType
from backend.review import run_review
from backend.writeback import build_writeback_comment, MockApprovalSystem

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEST_DOCX = os.path.join(ROOT, "测试合同_混合样例.docx")

CONTRACT_BODY = """甲方：测试科技有限公司
乙方：演示供应商有限公司
统一社会信用代码：91110108MA01TESTX0
合同金额：人民币 120,000 元
履行期限：自合同签订之日起 60 日
生效条件：双方盖章并签字后生效

第一条 标的物条款
甲方向乙方采购测试用服务器一台及配套定制软件。

第二条 知识产权条款
本合同所涉定制软件之全部知识产权归乙方所有，甲方仅获得非独占使用许可。

第三条 付款进度
甲方应在设备到货后 3 个工作日内，向乙方支付合同全部款项人民币 120,000 元。

第四条 违约责任
任一方违约，向守约方支付违约金，违约金不超过合同总额 20%。

第五条 保密义务
双方对本合同知悉的商业秘密负有保密义务。

第六条 争议解决管辖
因本合同发生争议，提交乙方所在地有管辖权的人民法院诉讼解决。
"""


def make_test_docx():
    if os.path.exists(TEST_DOCX):
        return
    doc = Document()
    doc.add_heading("测试采购合同（混合风险样例）", level=0)
    for line in CONTRACT_BODY.split("\n"):
        doc.add_paragraph(line)
    doc.save(TEST_DOCX)
    print(f"[生成] 测试合同: {TEST_DOCX}")


def main():
    make_test_docx()

    # 建任务并跑全链路
    task = store.create_task(contract_name="测试采购合同（混合风险样例）",
                             biz_type=BizType.PROCUREMENT.value, applicant="测试员")
    run_review(task, store.RULES, TEST_DOCX)

    print("\n===== 自测结果 =====")
    print(f"任务状态      : {task.status}")
    print(f"综合风险等级  : {task.risk_level}")
    print(f"识别风险数    : {len(task.risks)}")
    for r in task.risks:
        print(f"  - [{r.level.value}] {r.title}  (高亮定位={r.highlight})")

    # 断言
    assert task.status.value == "completed", "审查应完成"
    assert task.doc_text.strip(), "应解析出正文"
    assert task.segments, "应切分出条款"
    assert task.meta.amount == 120000.0, f"金额应提取为 120000，实际 {task.meta.amount}"
    assert task.meta.credit_code, "应提取统一社会信用代码"
    titles = {r.title for r in task.risks}
    assert "知识产权归属不利" in titles, "应命中知识产权归属不利"
    assert "付款无前置验收条件" in titles, "应命中付款无前置验收"
    assert any(r.level.value == "high" for r in task.risks), "应存在高风险"
    assert build_writeback_comment(task).strip(), "应生成回写评论文本"
    MockApprovalSystem().post_comment(task)  # 模拟回写不报错
    print("\n✅ PASS：解析 / 规则审查 / 元数据提取 / 回写文本生成 全部通过")

    print("\n提示：该 .docx 也可在系统「审查大盘」页用「上传并审查」手动上传验证 UI。")


if __name__ == "__main__":
    try:
        main()
    except AssertionError as e:
        print(f"\n❌ FAIL：{e}")
        sys.exit(1)
    except Exception:
        print("\n❌ ERROR：")
        traceback.print_exc()
        sys.exit(1)
